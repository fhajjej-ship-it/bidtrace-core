# SPDX-License-Identifier: MPL-2.0
"""Extract original bytes plus honest page/sheet coverage and immutable previews."""
from __future__ import annotations

from io import BytesIO
import ctypes
from pathlib import PurePosixPath
import hashlib
import re
import shutil
import subprocess
import tempfile
import time
import textwrap
import zipfile

from threading import RLock

from pypdf import PdfReader
import pypdfium2 as pdfium
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

MAX_BYTES = 25 * 1024 * 1024
MAX_PAGES = 300
MAX_TEXT = 2_000_000
PDFIUM_LOCK = RLock()  # PDFium calls must not overlap across upload threads.
MIME_TYPES = {
    '.pdf': 'application/pdf', '.txt': 'text/plain', '.md': 'text/markdown',
    '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
}


class ExtractionError(ValueError):
    pass


def _pdf_page_text(page):
    """Order PDFium's actual words by page geometry, including form XObjects.

    Text values come only from the PDF text layer. We never complete words from
    a dictionary or from benchmark references. PDFium supplies word boundaries;
    geometry supplies reading order and indentation, not new content.
    """
    textpage = page.get_textpage()
    try:
        words, chars, boxes, baselines = [], [], [], []

        def finish_word():
            if chars and boxes:
                left=min(b[0] for b in boxes);bottom=min(b[1] for b in boxes)
                right=max(b[2] for b in boxes);top=max(b[3] for b in boxes)
                words.append({'text': ''.join(chars), 'left':left, 'right':right,
                              'baseline':sorted(baselines)[len(baselines)//2], 'height':top-bottom})
            chars.clear();boxes.clear();baselines.clear()

        for index in range(textpage.count_chars()):
            value=pdfium.raw.FPDFText_GetUnicode(textpage.raw,index)
            char=chr(value) if value else ''
            if not char or char.isspace():
                finish_word()
                continue
            box=textpage.get_charbox(index,loose=True)
            origin_x,origin_y=ctypes.c_double(),ctypes.c_double()
            if not pdfium.raw.FPDFText_GetCharOrigin(textpage.raw,index,ctypes.byref(origin_x),ctypes.byref(origin_y)):
                raise ExtractionError('PDF text position could not be read completely.')
            if boxes:
                previous=boxes[-1]
                gap=box[0]-previous[2]
                height=min(box[3]-box[1],previous[3]-previous[1])
                if gap > max(1.0,.1*height) or abs((box[1]+box[3]-previous[1]-previous[3])/2)>max(2,.3*height):
                    finish_word()
            chars.append(char)
            boxes.append(box)
            baselines.append(origin_y.value)
        finish_word()
        rows=[]
        for word in sorted(words,key=lambda w:(-w['baseline'],w['left'])):
            row=next((r for r in rows if abs(r['baseline']-word['baseline']) <= 1.5),None)
            if row is None:
                row={'baseline':word['baseline'],'height':word['height'],'words':[]};rows.append(row)
            row['words'].append(word)
        lines=[];previous_row=None
        for row in sorted(rows,key=lambda r:-r['baseline']):
            ordered=sorted(row['words'],key=lambda w:w['left'])
            widths=[(w['right']-w['left'])/max(1,len(w['text'])) for w in ordered if w['right']>w['left']]
            width=sorted(widths)[len(widths)//2] if widths else 6
            if previous_row and previous_row['baseline']-row['baseline'] > 1.35*max(previous_row['height'],row['height']):
                lines.append('')
            parts=[' '*min(100,round(ordered[0]['left']/max(width,1)))];previous=None
            for word in ordered:
                if previous is not None:
                    gap=word['left']-previous['right']
                    parts.append(' '*max(1,min(80,round(gap/max(width,1)))))
                parts.append(word['text']);previous=word
            lines.append(''.join(parts))
            previous_row=row
        return '\n'.join(lines).strip()
    finally:
        textpage.close()


def _office_archive(content):
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024:
                raise ExtractionError('Office archive expands beyond the supported extraction limit.')
    except zipfile.BadZipFile as exc:
        raise ExtractionError('This is not a readable DOCX/XLSX file.') from exc


def _text_segments(text):
    """Deterministic generated pages; these are not claimed as original pagination."""
    pages = []
    for source_page in text.split('\f'):
        lines = []
        for line in source_page.split('\n'):
            lines.extend(textwrap.wrap(line, width=88, replace_whitespace=False,
                                       drop_whitespace=False) or [''])
        for start in range(0, len(lines), 48):
            pages.append('\n'.join(lines[start:start + 48]))
    return pages or ['']


def _render_text(text, label):
    image = Image.new('RGB', (744, 1053), 'white')
    draw = ImageDraw.Draw(image)
    # Use the operating system font; no additional font files are distributed.
    try:
        font = ImageFont.truetype('/System/Library/Fonts/Menlo.ttc', 12)
    except OSError:
        font = ImageFont.load_default(size=12)
    draw.text((40, 30), label, fill='#666666', font=font)
    bounds = draw.multiline_textbbox((40, 58), text, font=font, spacing=4)
    if bounds[2] > 710 or bounds[3] > 1010:
        raise ExtractionError('Generated page could not render completely.')
    draw.multiline_text((40, 58), text, fill='#222222', font=font, spacing=4)
    result = BytesIO()
    image.save(result, format='PNG')
    return result.getvalue()


def extract_document(filename, content, mime_type, object_store, prefix):
    if not isinstance(content, bytes) or not content:
        raise ExtractionError('The document is empty.')
    if len(content) > MAX_BYTES:
        raise ExtractionError('The document exceeds the 25 MB extraction limit.')
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix not in MIME_TYPES:
        raise ExtractionError('Supported files are PDF, TXT, MD, DOCX, and XLSX.')
    if not re.fullmatch(r'[A-Za-z0-9_\-/]+', prefix) or any(p in {'', '.', '..'} for p in prefix.split('/')):
        raise ExtractionError('Invalid storage prefix.')
    detected_mime = MIME_TYPES[suffix]
    raw_sha256 = hashlib.sha256(content).hexdigest()
    original_key = f'{prefix}/original/{raw_sha256}/source{suffix}'
    pages = []
    images = []
    ocr_started = time.monotonic()
    ocr_binary = shutil.which("tesseract")
    ocr_pages = []
    coverage = {'format': suffix[1:], 'original_bytes_preserved': True, 'warnings': [],
                'complete_text': True, 'page_numbering': 'original_pdf' if suffix == '.pdf' else 'generated_preview'}
    try:
        if suffix == '.pdf':
            if not content.startswith(b'%PDF-'):
                raise ExtractionError('The .pdf file does not contain a PDF header.')
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted and not reader.decrypt(''):
                raise ExtractionError('This PDF is password protected. Upload an unlocked copy.')
            with PDFIUM_LOCK, pdfium.PdfDocument(content) as pdf:
                if not 1 <= len(pdf) <= MAX_PAGES:
                    raise ExtractionError('PDF page count exceeds the supported 300-page limit or is empty.')
                for number, page in enumerate(pdf, 1):
                    page_text = _pdf_page_text(page)
                    page_record = {'page': number, 'text': page_text, 'image_kind': 'original_pdf_render',
                                   'text_method': 'pdf_text_layer' if page_text else 'none'}
                    # Bound each render for unusually large engineering drawings.
                    scale = min(1.5, 1800 / max(page.get_size()))
                    bitmap = page.render(scale=scale)
                    try:
                        rendered = BytesIO()
                        bitmap.to_pil().save(rendered, format='PNG')
                        png = rendered.getvalue()
                    finally:
                        bitmap.close()
                    images.append(png)
                    if not page_text and ocr_binary and time.monotonic() - ocr_started < 90:
                        try:
                            with tempfile.NamedTemporaryFile(suffix='.png') as preview:
                                preview.write(png)
                                preview.flush()
                                ocr = subprocess.run([ocr_binary, preview.name, 'stdout', '-l', 'eng'],
                                    capture_output=True, timeout=min(30, max(1, 90 - (time.monotonic() - ocr_started))), check=True)
                            recognized = ocr.stdout.decode('utf-8').strip()
                            if recognized:
                                page_record['text'] = recognized
                                page_record['text_method'] = 'ocr'
                                ocr_pages.append(number)
                        except (subprocess.SubprocessError, OSError, UnicodeError):
                            coverage['warnings'].append(f'OCR did not complete for PDF page {number}; its original image is retained.')
                    pages.append(page_record)
            image_only = [p['page'] for p in pages if not p['text']]
            coverage['image_only_pages'] = image_only
            coverage['ocr_pages'] = ocr_pages
            coverage['ocr_engine'] = 'tesseract-eng' if ocr_pages else None
            if ocr_pages:
                coverage['complete_text'] = False
                coverage['warnings'].append('OCR text was extracted locally from scanned pages; a human must compare quoted claims with the original page image before confirming them.')
            if image_only:
                coverage['complete_text'] = False
                coverage['warnings'].append('Some PDF pages have no text layer; their original page images are preserved. OCR was unavailable, empty, failed, or exceeded its bounded time budget.')
        elif suffix in {'.txt', '.md'}:
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError as exc:
                raise ExtractionError('TXT/MD uploads must use UTF-8 encoding.') from exc
            if '\x00' in text:
                raise ExtractionError('This text file contains binary data.')
            pages = [{'page': n, 'text': part, 'image_kind': 'generated_text_preview'}
                     for n, part in enumerate(_text_segments(text), 1)]
            coverage['warnings'].append('Images are generated previews of original text, not original PDF pages.')
        elif suffix == '.docx':
            _office_archive(content)
            doc = Document(BytesIO(content))
            blocks = []
            for node in doc.element.body:
                if node.tag.endswith('}p'):
                    blocks.append(Paragraph(node, doc).text)
                elif node.tag.endswith('}tbl'):
                    for row in Table(node, doc).rows:
                        blocks.append(' | '.join(cell.text for cell in row.cells))
            text = '\n'.join(blocks)
            if not text.strip():
                raise ExtractionError('DOCX has no readable paragraphs or table cells.')
            pages = [{'page': n, 'text': part, 'image_kind': 'generated_docx_text_preview'}
                     for n, part in enumerate(_text_segments(text), 1)]
            coverage['warnings'].append('DOCX pages are generated paragraph/table previews. Original layout, headers, footers, text boxes, and embedded images are not interpreted.')
            coverage['complete_text'] = False
        else:
            _office_archive(content)
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            sheets = []
            try:
                for sheet in workbook.worksheets:
                    if (sheet.max_row or 0) * (sheet.max_column or 0) > 500_000 or (sheet.max_row or 0) > 20_000:
                        raise ExtractionError('A spreadsheet sheet exceeds the supported 500,000-cell / 20,000-row limit.')
                    lines = [f'Sheet: {sheet.title}']
                    formulas = 0
                    for row in sheet.iter_rows():
                        values = []
                        for cell in row:
                            if cell.value is not None:
                                formulas += int(cell.data_type == 'f')
                                values.append(f'{cell.coordinate}: {cell.value}')
                        if values:
                            lines.append(' | '.join(values))
                    segments = _text_segments('\n'.join(lines))
                    sheet_pages = []
                    for segment in segments:
                        page_number = len(pages) + 1
                        sheet_pages.append(page_number)
                        pages.append({'page': page_number, 'text': segment, 'sheet': sheet.title,
                                      'image_kind': 'generated_sheet_preview'})
                    sheets.append({'name': sheet.title, 'pages': sheet_pages, 'state': sheet.sheet_state,
                                   'formula_count': formulas})
            finally:
                workbook.close()
            coverage['sheets'] = sheets
            coverage['page_numbering'] = 'generated_sheet_pages'
            coverage['warnings'].append('Sheet previews preserve names and cell references. Formulas are expressions, not recalculated values; charts, styles, and embedded images are not interpreted.')
            coverage['complete_text'] = False
        if not pages or len(pages) > MAX_PAGES:
            raise ExtractionError('Extraction exceeds the supported 300 generated pages.')
        text = '\f'.join(page['text'] for page in pages)
        if len(text) > MAX_TEXT:
            raise ExtractionError('Extracted text exceeds the 2 million character limit.')
        if suffix != '.pdf':
            images = [_render_text(page['text'], f'{filename} · Generated preview {page["page"]}') for page in pages]
        # Persist only after the entire extraction succeeded; bytes never overwrite another snapshot.
        object_store.put(original_key, content, detected_mime)
        for page, png in zip(pages, images, strict=True):
            digest = hashlib.sha256(png).hexdigest()
            page['image_key'] = f'{prefix}/pages/{digest}/page.png'
            object_store.put(page['image_key'], png, 'image/png')
        coverage['pages'] = len(pages)
        coverage['text_characters'] = len(text)
        return {'text': text, 'pages': pages, 'mime_type': detected_mime, 'object_key': original_key,
                'raw_sha256': raw_sha256, 'coverage': coverage}
    except ExtractionError:
        raise
    except Exception as exc:
        # Keep storage/service failures distinguishable from corrupt document content.
        if exc.__class__.__module__.startswith(('botocore', 'boto3')):
            raise
        raise ExtractionError(f'Could not extract {suffix[1:].upper()} document: {type(exc).__name__}.') from exc
